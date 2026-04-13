from __future__ import annotations

import logging
from datetime import datetime, timedelta
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ============================================================
# PATHS
# ============================================================

BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "input"
OUTPUT_DIR = BASE_DIR / "output"
LOG_DIR = BASE_DIR / "log"

INPUT_FILE = INPUT_DIR / "invoice_data.csv"
LOG_FILE = LOG_DIR / "run.log"

# ============================================================
# CONFIG
# ============================================================

COMPANY_NAME = "Northstar Automation Services"
COMPANY_EMAIL = "billing@northstar-automation.com"
COMPANY_PHONE = "+61 4XX XXX XXX"
COMPANY_ABN = "ABN 12 345 678 901"
PAYMENT_TERMS_DAYS = 14
DEFAULT_CURRENCY_SYMBOL = "$"

REQUIRED_COLUMNS = [
    "invoice_id",
    "client_name",
    "client_email",
    "invoice_date",
    "item_description",
    "quantity",
    "unit_price",
    "tax_rate",
]

# ============================================================
# LOGGING
# ============================================================

def setup_logging() -> None:
    LOG_DIR.mkdir(parents=True, exist_ok=True)

    logger = logging.getLogger()
    logger.setLevel(logging.INFO)

    if logger.handlers:
        logger.handlers.clear()

    formatter = logging.Formatter(
        "%(asctime)s | %(levelname)s | %(message)s"
    )

    file_handler = logging.FileHandler(LOG_FILE, encoding="utf-8")
    file_handler.setFormatter(formatter)

    stream_handler = logging.StreamHandler()
    stream_handler.setFormatter(formatter)

    logger.addHandler(file_handler)
    logger.addHandler(stream_handler)


# ============================================================
# DATA LOADING
# ============================================================

def load_data(file_path: Path) -> pd.DataFrame:
    logging.info("Loading input data from: %s", file_path)

    if not file_path.exists():
        raise FileNotFoundError(
            f"Input file not found: {file_path}\n"
            f"Create the file and place it in the input folder."
        )

    if file_path.suffix.lower() == ".csv":
        df = pd.read_csv(file_path)
    elif file_path.suffix.lower() in {".xlsx", ".xls"}:
        df = pd.read_excel(file_path)
    else:
        raise ValueError("Unsupported input format. Use CSV or Excel.")

    df = df.dropna(how="all")
    logging.info("Loaded %s rows", len(df))
    return df


# ============================================================
# DATA VALIDATION AND CLEANING
# ============================================================

def standardize_columns(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.columns = (
        df.columns.astype(str)
        .str.strip()
        .str.lower()
        .str.replace(" ", "_")
        .str.replace("-", "_")
    )
    return df


def validate_required_columns(df: pd.DataFrame, required_columns: list[str]) -> None:
    missing = [col for col in required_columns if col not in df.columns]
    if missing:
        raise ValueError(
            f"Missing required columns: {missing}\n"
            f"Expected columns: {required_columns}"
        )


def clean_data(df: pd.DataFrame) -> pd.DataFrame:
    logging.info("Cleaning input data")
    df = standardize_columns(df)
    validate_required_columns(df, REQUIRED_COLUMNS)

    df = df.copy()

    text_columns = ["invoice_id", "client_name", "client_email", "item_description"]
    for col in text_columns:
        df[col] = df[col].astype(str).str.strip()

    df = df[df["invoice_id"] != ""]
    df = df[df["client_name"] != ""]
    df = df[df["item_description"] != ""]

    df["invoice_date"] = pd.to_datetime(df["invoice_date"], errors="coerce")
    if df["invoice_date"].isna().any():
        bad_rows = df[df["invoice_date"].isna()].index.tolist()
        raise ValueError(f"Invalid invoice_date values found in rows: {bad_rows}")

    numeric_columns = ["quantity", "unit_price", "tax_rate"]
    for col in numeric_columns:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    if df[numeric_columns].isna().any().any():
        bad_mask = df[numeric_columns].isna().any(axis=1)
        bad_rows = df[bad_mask].index.tolist()
        raise ValueError(f"Invalid numeric values found in rows: {bad_rows}")

    if (df["quantity"] <= 0).any():
        bad_rows = df[df["quantity"] <= 0].index.tolist()
        raise ValueError(f"Quantity must be greater than 0. Bad rows: {bad_rows}")

    if (df["unit_price"] < 0).any():
        bad_rows = df[df["unit_price"] < 0].index.tolist()
        raise ValueError(f"Unit price cannot be negative. Bad rows: {bad_rows}")

    if ((df["tax_rate"] < 0) | (df["tax_rate"] > 1)).any():
        bad_rows = df[((df["tax_rate"] < 0) | (df["tax_rate"] > 1))].index.tolist()
        raise ValueError(
            f"tax_rate must be between 0 and 1, e.g. 0.10 for 10%. Bad rows: {bad_rows}"
        )

    df["line_subtotal"] = df["quantity"] * df["unit_price"]
    df["line_tax"] = df["line_subtotal"] * df["tax_rate"]
    df["line_total"] = df["line_subtotal"] + df["line_tax"]
    df["due_date"] = df["invoice_date"] + timedelta(days=PAYMENT_TERMS_DAYS)

    logging.info("Data cleaned successfully")
    return df


# ============================================================
# EXCEL FORMATTING HELPERS
# ============================================================

def apply_currency(cell) -> None:
    cell.number_format = f'"{DEFAULT_CURRENCY_SYMBOL}"#,##0.00'


def apply_date(cell) -> None:
    cell.number_format = "yyyy-mm-dd"


def auto_fit_columns(ws) -> None:
    for column_cells in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)

        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            if len(value) > max_length:
                max_length = len(value)

        ws.column_dimensions[column_letter].width = min(max_length + 2, 35)


def style_table_header(row: Iterable) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    alignment = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in row:
        cell.fill = fill
        cell.font = font
        cell.alignment = alignment
        cell.border = border


def style_cells_with_border(cells: Iterable) -> None:
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in cells:
        cell.border = border
        cell.alignment = Alignment(vertical="center")


# ============================================================
# INVOICE GENERATION
# ============================================================

def create_invoice_workbook(invoice_df: pd.DataFrame) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"

    invoice_id = invoice_df["invoice_id"].iloc[0]
    client_name = invoice_df["client_name"].iloc[0]
    client_email = invoice_df["client_email"].iloc[0]
    invoice_date = invoice_df["invoice_date"].iloc[0]
    due_date = invoice_df["due_date"].iloc[0]

    subtotal = float(invoice_df["line_subtotal"].sum())
    tax_total = float(invoice_df["line_tax"].sum())
    grand_total = float(invoice_df["line_total"].sum())

    # Top section
    ws["A1"] = COMPANY_NAME
    ws["A1"].font = Font(size=16, bold=True)

    ws["A2"] = COMPANY_EMAIL
    ws["A3"] = COMPANY_PHONE
    ws["A4"] = COMPANY_ABN

    ws["F1"] = "INVOICE"
    ws["F1"].font = Font(size=18, bold=True)
    ws["F1"].alignment = Alignment(horizontal="right")

    ws["F2"] = "Invoice ID:"
    ws["G2"] = invoice_id
    ws["F3"] = "Invoice Date:"
    ws["G3"] = invoice_date
    ws["F4"] = "Due Date:"
    ws["G4"] = due_date

    apply_date(ws["G3"])
    apply_date(ws["G4"])

    # Bill to
    ws["A6"] = "Bill To"
    ws["A6"].font = Font(bold=True, size=12)

    ws["A7"] = client_name
    ws["A8"] = client_email

    # Table
    table_start_row = 10
    headers = [
        "Description",
        "Quantity",
        "Unit Price",
        "Tax Rate",
        "Subtotal",
        "Tax",
        "Line Total",
    ]

    for col_num, header in enumerate(headers, start=1):
        ws.cell(row=table_start_row, column=col_num, value=header)

    style_table_header(ws[table_start_row])

    current_row = table_start_row + 1

    for _, row in invoice_df.iterrows():
        ws.cell(current_row, 1, row["item_description"])
        ws.cell(current_row, 2, row["quantity"])
        ws.cell(current_row, 3, row["unit_price"])
        ws.cell(current_row, 4, row["tax_rate"])
        ws.cell(current_row, 5, row["line_subtotal"])
        ws.cell(current_row, 6, row["line_tax"])
        ws.cell(current_row, 7, row["line_total"])

        style_cells_with_border(ws[current_row])

        ws.cell(current_row, 3).number_format = '"$"#,##0.00'
        ws.cell(current_row, 4).number_format = "0.00%"
        ws.cell(current_row, 5).number_format = '"$"#,##0.00'
        ws.cell(current_row, 6).number_format = '"$"#,##0.00'
        ws.cell(current_row, 7).number_format = '"$"#,##0.00'

        current_row += 1

    # Summary block
    summary_start = current_row + 1

    ws.cell(summary_start, 6, "Subtotal").font = Font(bold=True)
    ws.cell(summary_start, 7, subtotal)
    apply_currency(ws.cell(summary_start, 7))

    ws.cell(summary_start + 1, 6, "Tax").font = Font(bold=True)
    ws.cell(summary_start + 1, 7, tax_total)
    apply_currency(ws.cell(summary_start + 1, 7))

    ws.cell(summary_start + 2, 6, "Total").font = Font(bold=True, size=12)
    ws.cell(summary_start + 2, 7, grand_total)
    ws.cell(summary_start + 2, 7).font = Font(bold=True, size=12)
    apply_currency(ws.cell(summary_start + 2, 7))

    for r in range(summary_start, summary_start + 3):
        style_cells_with_border(ws[r][5:7])

    # Notes
    note_row = summary_start + 5
    ws.cell(note_row, 1, "Payment Terms").font = Font(bold=True)
    ws.cell(
        note_row + 1,
        1,
        f"Payment due within {PAYMENT_TERMS_DAYS} days from the invoice date."
    )

    # Freeze pane
    ws.freeze_panes = "A11"

    # Alignment tweaks
    for row in ws.iter_rows():
        for cell in row:
            if cell.column in [2, 3, 4, 5, 6, 7]:
                cell.alignment = Alignment(horizontal="right", vertical="center")

    auto_fit_columns(ws)

    return wb


def save_invoice_file(wb: Workbook, invoice_id: str, client_name: str) -> Path:
    safe_client_name = "".join(
        char if char.isalnum() or char in (" ", "_", "-") else "_"
        for char in client_name
    ).strip().replace(" ", "_")

    filename = f"invoice_{invoice_id}_{safe_client_name}.xlsx"
    output_path = OUTPUT_DIR / filename
    wb.save(output_path)
    return output_path


def create_single_invoice(invoice_df: pd.DataFrame) -> Path:
    invoice_id = invoice_df["invoice_id"].iloc[0]
    client_name = invoice_df["client_name"].iloc[0]

    wb = create_invoice_workbook(invoice_df)
    output_path = save_invoice_file(wb, invoice_id, client_name)

    logging.info("Created invoice file: %s", output_path.name)
    return output_path


def process_all_invoices(df: pd.DataFrame) -> list[Path]:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    created_files: list[Path] = []

    grouped = df.groupby("invoice_id", sort=True)

    logging.info("Processing %s invoices", len(grouped))

    for invoice_id, invoice_df in grouped:
        logging.info("Generating invoice: %s", invoice_id)
        created_file = create_single_invoice(invoice_df.reset_index(drop=True))
        created_files.append(created_file)

    return created_files


# ============================================================
# SUMMARY REPORT
# ============================================================

def create_run_summary(df: pd.DataFrame, created_files: list[Path]) -> Path:
    summary = (
        df.groupby(["invoice_id", "client_name", "client_email"], as_index=False)
        .agg(
            invoice_date=("invoice_date", "min"),
            due_date=("due_date", "min"),
            subtotal=("line_subtotal", "sum"),
            tax=("line_tax", "sum"),
            total=("line_total", "sum"),
            line_count=("item_description", "count"),
        )
        .sort_values("invoice_id")
    )

    summary["generated_file"] = [
        file.name for file in created_files
    ]

    output_file = OUTPUT_DIR / "invoice_run_summary.xlsx"
    summary.to_excel(output_file, index=False)

    logging.info("Created run summary: %s", output_file.name)
    return output_file


# ============================================================
# MAIN
# ============================================================

def main() -> None:
    setup_logging()
    logging.info("Starting Invoice Automation System")

    try:
        INPUT_DIR.mkdir(parents=True, exist_ok=True)
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        LOG_DIR.mkdir(parents=True, exist_ok=True)

        df = load_data(INPUT_FILE)
        cleaned_df = clean_data(df)
        created_files = process_all_invoices(cleaned_df)
        create_run_summary(cleaned_df, created_files)

        logging.info("Completed successfully. %s invoice files created.", len(created_files))

    except Exception as exc:
        logging.exception("Process failed: %s", exc)
        raise


if __name__ == "__main__":
    main()